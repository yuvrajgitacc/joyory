import os
from pptx import Presentation
from pptx.util import Inches

def add_workflow_slide(source_pptx, output_pptx, image_path):
    prs = Presentation(source_pptx)
    blank_slide_layout = prs.slide_layouts[6] # completely blank layout
    
    # Add new slide
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Add the image full bleed 20.0 x 11.25 inches
    left = Inches(0)
    top = Inches(0)
    width = prs.slide_width
    height = prs.slide_height
    
    slide.shapes.add_picture(image_path, left, top, width, height)
    
    # Move the new slide to be right after Slide 5 (index 5, making it slide 6)
    xml_slides = prs.slides._sldIdLst
    new_slide_elem = xml_slides[-1]
    # Insert at position 5 (0-indexed, so 6th slide)
    xml_slides.remove(new_slide_elem)
    xml_slides.insert(5, new_slide_elem)
    
    prs.save(output_pptx)
    print(f"Successfully created: {output_pptx} with new workflow slide at position 6 (Total slides: {len(prs.slides)})")

if __name__ == "__main__":
    src = "Joyory-AI-Beauty-Concierge .pptx"
    out = "Joyory-AI-Beauty-Concierge-With-Workflow.pptx"
    img = "joyory_workflow_dark.png"
    add_workflow_slide(src, out, img)
